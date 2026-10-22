"""Word "fields": table of contents and page numbers.

python-docx has no built-in support for these, so we write the underlying
XML ourselves. A Word *field* is a small instruction embedded in the text,
such as "show the current page number" or "build a table of contents from the
headings". Word evaluates the instruction when it displays the document.

A field is stored as a sequence of runs:

    [begin marker] [instruction text] [separator] [result text] [end marker]

The "result text" is what is shown until Word recalculates the field. We fill
it in ourselves so the document looks right even before Word updates it.
"""

from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt


def _add_field_char(paragraph, kind):
    """Append a begin / separate / end marker run to the paragraph."""

    run = paragraph.add_run()
    marker = OxmlElement("w:fldChar")
    marker.set(qn("w:fldCharType"), kind)
    run._r.append(marker)

    return run


def _add_instruction(paragraph, instruction):
    """Append the field's instruction text, e.g. ' PAGE '."""

    run = paragraph.add_run()
    element = OxmlElement("w:instrText")
    element.set(qn("xml:space"), "preserve")
    element.text = instruction
    run._r.append(element)

    return run


def add_page_number_footer(document):
    """Show a centred 'Page N' in the footer of every page."""

    for section in document.sections:
        paragraph = section.footer.paragraphs[0]
        paragraph.text = ""
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

        paragraph.add_run("Page ")
        _add_field_char(paragraph, "begin")
        _add_instruction(paragraph, " PAGE ")
        _add_field_char(paragraph, "separate")
        paragraph.add_run("1")  # placeholder; Word/LibreOffice replace it
        _add_field_char(paragraph, "end")


def add_table_of_contents(document, headings, heading_size=14):
    """Insert a table of contents listing `headings`, then a page break.

    `headings` is the list of section heading texts. The entries are written
    out directly so the contents list is visible immediately. The surrounding
    TOC field also lets Word rebuild the list with page numbers: after
    enable_update_fields_on_open(), Word offers to do so when the file opens.
    """

    title = document.add_paragraph()
    title_run = title.add_run("Table of Contents")
    title_run.bold = True
    title_run.font.size = Pt(heading_size)

    entries = list(headings) or [""]
    last_index = len(entries) - 1

    for index, text in enumerate(entries):
        paragraph = document.add_paragraph()

        if index == 0:
            _add_field_char(paragraph, "begin")
            _add_instruction(paragraph, ' TOC \\o "1-3" \\h \\z \\u ')
            _add_field_char(paragraph, "separate")

        paragraph.add_run(text)

        if index == last_index:
            _add_field_char(paragraph, "end")

    # Start the report body on a fresh page.
    document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def enable_update_fields_on_open(document):
    """Ask Word to refresh all fields (TOC page numbers) when the file opens.

    Word will show a one-time prompt asking whether to update the fields.
    """

    settings = document.settings.element

    if settings.find(qn("w:updateFields")) is None:
        element = OxmlElement("w:updateFields")
        element.set(qn("w:val"), "true")
        settings.append(element)
