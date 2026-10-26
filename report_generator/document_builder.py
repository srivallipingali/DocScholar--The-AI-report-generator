
import datetime
from pathlib import Path
from io import BytesIO

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image


def _hex_color(value):
    """Convert a #RRGGBB template color into Word RGBColor safely."""

    value = str(value or "#6A1B4D").strip().lstrip("#")
    if len(value) != 6 or any(char not in "0123456789abcdefABCDEF" for char in value):
        value = "6A1B4D"
    return RGBColor(int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))


def _prepare_output(output_path):
    """Create the output folder if a file path was given.

    The app passes an in-memory BytesIO instead of a path, so only real
    paths need a folder.
    """

    if isinstance(output_path, (str, Path)):
        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

    return output_path


def _create_styled_document(config, template_bytes=None):
    """Create a report document, optionally using an uploaded DOCX as its base.

    When template_bytes is supplied, keep the uploaded document's styles,
    page setup, headers and footers, but remove its old body content so the
    newly generated report does not mix with the previous report.
    """

    if template_bytes:
        document = Document(BytesIO(template_bytes))
        body = document._element.body
        section_properties = body.sectPr
        for child in list(body):
            if child is not section_properties:
                body.remove(child)
        # The template already owns its margins, styles and theme. Do not
        # overwrite those with the built-in template configuration.
        return document

    document = Document()

    # Read optional template colors. Existing templates keep their original look.
    accent = _hex_color(config.get("accent_color", "#6A1B4D"))
    text_color = _hex_color(config.get("text_color", "#29232A"))

    # Page margins.
    margins = config["margins"]

    for section in document.sections:
        section.top_margin = Inches(margins["top"])
        section.bottom_margin = Inches(margins["bottom"])
        section.left_margin = Inches(margins["left"])
        section.right_margin = Inches(margins["right"])

    # Default font for normal paragraphs.
    font_name = config["font_name"]
    font_size = config["font_size"]

    normal_style = document.styles["Normal"]
    normal_style.font.name = font_name
    normal_style.font.size = Pt(font_size)
    normal_style.font.color.rgb = text_color

    # Main report title style.
    title_style = document.styles["Title"]
    title_style.font.name = config.get("title_font_name", font_name)
    title_style.font.size = Pt(
        config.get("title_font_size", 18)
    )
    title_style.font.bold = True
    title_style.font.color.rgb = accent

    # Section heading style.
    heading_style = document.styles["Heading 1"]
    heading_style.font.name = config.get("heading_font_name", font_name)
    heading_style.font.size = Pt(
        config.get("heading_font_size", 14)
    )
    heading_style.font.bold = True
    heading_style.font.color.rgb = accent

    return document


def _append_field(paragraph, field_name):
    """Add a dynamic Word field code (e.g. PAGE, NUMPAGES) properly wrapped in runs."""
    run_begin = paragraph.add_run()
    fld1 = OxmlElement("w:fldChar")
    fld1.set(qn("w:fldCharType"), "begin")
    run_begin._r.append(fld1)

    run_text = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_name
    run_text._r.append(instr)

    run_sep = paragraph.add_run()
    fld2 = OxmlElement("w:fldChar")
    fld2.set(qn("w:fldCharType"), "separate")
    run_sep._r.append(fld2)

    run_end = paragraph.add_run()
    fld3 = OxmlElement("w:fldChar")
    fld3.set(qn("w:fldCharType"), "end")
    run_end._r.append(fld3)


def _configure_page_numbering(document):
    """Enable proper dynamic page numbering (Page X of Y) in the footer for pages after the cover."""
    for section in document.sections:
        section.different_first_page_header_footer = True
        footer = section.footer
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        p.paragraph_format.space_before = Pt(8)

        run_page = p.add_run("Page ")
        run_page.font.size = Pt(9)
        run_page.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        _append_field(p, "PAGE")

        run_of = p.add_run(" of ")
        run_of.font.size = Pt(9)
        run_of.font.color.rgb = RGBColor(0x77, 0x77, 0x77)
        _append_field(p, "NUMPAGES")


def _add_cover_page(document, title, config, subtitle=None, abstract=None):
    """Generate an elegant, professional front cover page."""
    font_name = config.get("font_name", "Arial")
    accent = _hex_color(config.get("accent_color", "#6A1B4D"))
    template_name = config.get("template_name", "Report").upper()

    p_top = document.add_paragraph()
    p_top.paragraph_format.space_before = Pt(60)

    p_tag = document.add_paragraph()
    tag_run = p_tag.add_run(template_name)
    tag_run.font.name = font_name
    tag_run.font.size = Pt(10)
    tag_run.font.bold = True
    tag_run.font.color.rgb = accent
    p_tag.paragraph_format.space_after = Pt(12)

    p_title = document.add_paragraph()
    title_run = p_title.add_run(title)
    title_run.font.name = config.get("title_font_name", font_name)
    title_run.font.size = Pt(max(24, config.get("title_font_size", 18) + 8))
    title_run.font.bold = True
    title_run.font.color.rgb = accent
    p_title.paragraph_format.space_after = Pt(12)

    p_line = document.add_paragraph()
    line_run = p_line.add_run("―" * 32)
    line_run.font.color.rgb = accent
    line_run.font.size = Pt(11)
    p_line.paragraph_format.space_after = Pt(18)

    if subtitle:
        p_sub = document.add_paragraph()
        sub_run = p_sub.add_run(subtitle)
        sub_run.font.name = font_name
        sub_run.font.size = Pt(13)
        sub_run.font.italic = True
        sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        p_sub.paragraph_format.space_after = Pt(16)

    if abstract:
        p_abs = document.add_paragraph()
        abs_tag = p_abs.add_run("Executive Overview: ")
        abs_tag.font.name = font_name
        abs_tag.font.bold = True
        abs_tag.font.size = Pt(10.5)
        abs_run = p_abs.add_run(abstract)
        abs_run.font.name = font_name
        abs_run.font.size = Pt(10.5)
        abs_run.font.italic = True
        abs_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        p_abs.paragraph_format.space_after = Pt(24)

    p_space = document.add_paragraph()
    p_space.paragraph_format.space_after = Pt(80)

    author = config.get("author") or "DocScholar AI"
    date_str = config.get("date") or datetime.date.today().strftime("%B %d, %Y")

    p_meta = document.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(4)
    auth_label = p_meta.add_run("Prepared by: ")
    auth_label.font.name = font_name
    auth_label.font.bold = True
    auth_label.font.size = Pt(10.5)
    auth_val = p_meta.add_run(author)
    auth_val.font.name = font_name
    auth_val.font.size = Pt(10.5)

    p_date = document.add_paragraph()
    date_label = p_date.add_run("Date: ")
    date_label.font.name = font_name
    date_label.font.bold = True
    date_label.font.size = Pt(10.5)
    date_val = p_date.add_run(date_str)
    date_val.font.name = font_name
    date_val.font.size = Pt(10.5)

    document.add_page_break()


def _add_index_page(document, sections, config, start_page=3):
    """Generate a clean Table of Contents / Index page."""
    font_name = config.get("font_name", "Arial")
    accent = _hex_color(config.get("accent_color", "#6A1B4D"))
    margins = config.get("margins", {"left": 1.0, "right": 1.0})
    left_m = margins.get("left", 1.0)
    right_m = margins.get("right", 1.0)
    tab_position = Inches(max(4.0, 8.5 - left_m - right_m))

    toc_heading = document.add_heading("Table of Contents", level=1)
    toc_heading.paragraph_format.space_before = Pt(18)
    toc_heading.paragraph_format.space_after = Pt(6)

    p_caption = document.add_paragraph()
    cap_run = p_caption.add_run("Outline of Topics & Report Sections")
    cap_run.font.name = font_name
    cap_run.font.size = Pt(10)
    cap_run.font.italic = True
    cap_run.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    p_caption.paragraph_format.space_after = Pt(20)

    section_per_page = config.get("section_per_page", True)

    for i, section in enumerate(sections, 1):
        p = document.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(
            tab_position,
            WD_TAB_ALIGNMENT.RIGHT,
            WD_TAB_LEADER.DOTS,
        )
        p.paragraph_format.space_before = Pt(3)
        p.paragraph_format.space_after = Pt(8)

        num_run = p.add_run(f"{i:02d}  ")
        num_run.font.name = font_name
        num_run.font.bold = True
        num_run.font.size = Pt(11)
        num_run.font.color.rgb = accent

        heading_text = section.heading if hasattr(section, "heading") else str(section)
        heading_run = p.add_run(f"{heading_text}\t")
        heading_run.font.name = font_name
        heading_run.font.size = Pt(11)

        page_num = (start_page + i - 1) if section_per_page else start_page
        page_run = p.add_run(f"{page_num}")
        page_run.font.name = font_name
        page_run.font.bold = True
        page_run.font.size = Pt(10.5)

    document.add_page_break()


def _add_paragraphs(document, body):
    """Split body text on blank lines and add each block as a paragraph.

    Line breaks inside a block are joined with spaces so text copied from a
    PDF (which breaks lines mid-sentence) flows as normal paragraphs.
    """

    for paragraph in body.split("\n\n"):
        cleaned_paragraph = " ".join(
            line.strip()
            for line in paragraph.splitlines()
            if line.strip()
        )

        if cleaned_paragraph:
            document.add_paragraph(cleaned_paragraph)


def _fit_image_size(image_data, max_width=5.5, max_height=4.5):
    """Return page-friendly image dimensions while preserving aspect ratio."""
    try:
        with Image.open(BytesIO(image_data)) as image:
            width_px, height_px = image.size
        if not width_px or not height_px:
            raise ValueError("Image has no dimensions.")
        aspect_ratio = width_px / height_px
        width_inches = min(max_width, max_height * aspect_ratio)
        height_inches = width_inches / aspect_ratio
        return Inches(width_inches), Inches(height_inches)
    except Exception:
        # Let python-docx report invalid image data with the normal helpful error.
        return Inches(max_width), None


def _add_uploaded_images(document, images, start_index=1):
    """Append unplaced images with captions to the end of the report.

    Each image is a dictionary with a filename, image bytes, and an optional
    caption.
    """
    if not images:
        return

    document.add_page_break()
    document.add_heading("Uploaded Images", level=1)

    for index, image in enumerate(images, start=start_index):
        image_data = image.get("data")
        if not image_data:
            continue

        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            image_width, image_height = _fit_image_size(image_data, max_width=5.5, max_height=4.5)
            paragraph.add_run().add_picture(
                BytesIO(image_data),
                width=image_width,
                height=image_height,
            )
        except Exception as error:
            raise ValueError(
                f"Could not insert uploaded image "
                f"{image.get('name', f'Figure {index}')}: "
                "please upload a valid PNG or JPG image."
            ) from error

        caption_text = (image.get("caption") or "").strip()
        if not caption_text:
            caption_text = image.get("name") or f"Figure {index}"

        caption = document.add_paragraph(
            f"Figure {index}. {caption_text}",
            style="Caption",
        )
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        attribution = (image.get("attribution") or "").strip()
        if attribution and attribution != caption_text:
            attribution_paragraph = document.add_paragraph(
                f"Source: {attribution}",
                style="Caption",
            )
            attribution_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _add_images_inline(document, images, start_index=1):
    """Insert images at the current position with captions."""
    for index, image in enumerate(images, start=start_index):
        if not image.get("data"):
            continue
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        try:
            image_width, image_height = _fit_image_size(
                image["data"], max_width=5.5, max_height=4.5
            )
            paragraph.add_run().add_picture(
                BytesIO(image["data"]),
                width=image_width,
                height=image_height,
            )
        except Exception as error:
            raise ValueError(f"Could not insert image {image.get('name', index)}; use a valid image file.") from error
        caption_text = (image.get("caption") or "").strip() or image.get("name") or f"Figure {index}"
        caption = document.add_paragraph(f"Figure {index}. {caption_text}", style="Caption")
        caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
        attribution = (image.get("attribution") or "").strip()
        if attribution and attribution != caption_text:
            attribution_paragraph = document.add_paragraph(
                f"Source: {attribution}",
                style="Caption",
            )
            attribution_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER


def _save(document, output_path):
    """Save the document and return the same path (or BytesIO) it was given."""

    document.save(output_path)

    return output_path


def build_report(text, config, output_path, images=None, template_bytes=None):
    """
    Create a formatted Word report from plain text (no AI involved).

    The report title comes from the configuration; all input text is report
    content.
    """

    if not text or not text.strip():
        raise ValueError("Cannot generate a report from empty text.")

    title = (config.get("report_title") or "").strip()
    if not title:
        raise ValueError("A report title is required.")

    output_path = _prepare_output(output_path)
    document = _create_styled_document(config, template_bytes=template_bytes)

    include_cover = config.get("include_cover_page", False)

    if include_cover:
        _configure_page_numbering(document)

    if include_cover:
        _add_cover_page(document, title, config)
    else:
        document.add_paragraph(title, style="Title")

    _add_paragraphs(document, text.strip())
    _add_uploaded_images(document, images)

    return _save(document, output_path)


def build_structured_report(report, config, output_path, images=None, template_bytes=None):
    """
    Create a formatted Word report from a StructuredReport (the AI path).

    `report` is a report_schema.StructuredReport: a title plus a list of
    sections, each with a heading and content. Each section becomes a
    Heading 1 followed by its paragraphs.

    config["report_title"], if present and non-empty, overrides the title
    chosen by the AI, so a title typed by the user always wins.
    """

    if not report.sections:
        raise ValueError("Cannot generate a report with no sections.")

    output_path = _prepare_output(output_path)
    document = _create_styled_document(config, template_bytes=template_bytes)

    include_cover = config.get("include_cover_page", False)
    include_index = config.get("include_index", False)
    section_per_page = config.get("section_per_page", False)

    if include_cover or include_index:
        _configure_page_numbering(document)

    title = (config.get("report_title") or "").strip() or report.title
    subtitle = (config.get("report_subtitle") or "").strip() or getattr(report, "subtitle", None)
    abstract = (config.get("report_abstract") or "").strip() or getattr(report, "abstract", None)

    if include_cover:
        _add_cover_page(
            document,
            title,
            config,
            subtitle=subtitle,
            abstract=abstract,
        )
    else:
        document.add_paragraph(title, style="Title")

    if include_index:
        start_page = 3 if include_cover else 2
        _add_index_page(
            document,
            report.sections,
            config,
            start_page=start_page,
        )

    images_by_id = {str(image.get("id")): image for image in (images or []) if image.get("id") is not None}
    placed_ids = set()

    for i, section in enumerate(report.sections):
        if section_per_page and i > 0:
            document.add_page_break()
        document.add_heading(section.heading, level=1)
        _add_paragraphs(document, section.content)
        for image_id in getattr(section, "image_ids", []):
            image_id = str(image_id)
            image = images_by_id.get(image_id)
            if image_id in placed_ids:
                continue
            if image and image.get("data"):
                _add_images_inline(document, [image], start_index=len(placed_ids) + 1)
                placed_ids.add(image_id)

    # Unassigned images are preserved in a fallback section.
    unplaced = [
        image for image in (images or [])
        if image.get("data")
        and (image.get("id") is None or str(image.get("id")) not in placed_ids)
        # PPTX images are included only when Gemini explicitly selects them.
        and image.get("source") != "pptx"
    ]
    _add_uploaded_images(document, unplaced, start_index=len(placed_ids) + 1)

    return _save(document, output_path)
