from pathlib import Path

import pymupdf
from docx import Document


def read_file(file_path):
    """
    Read a TXT, PDF, or DOCX file and return its extracted text.
    """

    file_path = Path(file_path)

    # Step 1: Check whether the file exists.
    if not file_path.is_file():
        raise FileNotFoundError(
            f"File not found: {file_path}"
        )

    # Step 2: Identify the file type.
    extension = file_path.suffix.lower()

    # Step 3: Read a TXT file.
    if extension == ".txt":
        try:
            text = file_path.read_text(
                encoding="utf-8"
            ).strip()

        except UnicodeDecodeError as error:
            raise ValueError(
                "The TXT file is not encoded as UTF-8."
            ) from error

    # Step 4: Read a PDF file.
    elif extension == ".pdf":
        try:
            with pymupdf.open(file_path) as pdf:
                pages = [
                    page.get_text()
                    for page in pdf
                ]

            text = "\n".join(pages).strip()

        except Exception as error:
            raise ValueError(
                f"Could not read the PDF file: {error}"
            ) from error

    # Step 5: Read a DOCX file.
    elif extension == ".docx":
        try:
            document = Document(file_path)

            # Extract text from paragraphs.
            paragraphs = [
                paragraph.text
                for paragraph in document.paragraphs
            ]

            # Extract text from tables too.
            table_text = []

            for table in document.tables:
                for row in table.rows:
                    row_text = " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )

                    table_text.append(row_text)

            text_parts = paragraphs + table_text
            text = "\n".join(text_parts).strip()

        except Exception as error:
            raise ValueError(
                f"Could not read the DOCX file: {error}"
            ) from error

    # Step 6: Reject unsupported file types.
    else:
        raise ValueError(
            "Unsupported file type. Please use TXT, PDF, or DOCX."
        )

    # Step 7: Check that readable text was found.
    if not text:
        raise ValueError(
            "The file is empty or contains no readable text."
        )

    return text

