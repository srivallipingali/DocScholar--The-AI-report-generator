from pptx import Presentation
from pptx.util import Inches

from report_generator.pptx_reader import read_pptx


def test_read_pptx_preserves_slide_boundaries_and_image_context(tmp_path):
    from PIL import Image

    image_path = tmp_path / "diagram.png"
    Image.new("RGB", (80, 50), color="red").save(image_path)

    presentation = Presentation()
    slide1 = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide1.shapes.title.text = "DC Generator Construction"
    slide1.shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    ).text_frame.text = "The diagram shows the armature and field poles."
    slide1.shapes.add_picture(str(image_path), Inches(1), Inches(2), width=Inches(2))

    slide2 = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide2.shapes.title.text = "Working Principle"
    slide2.shapes.add_textbox(
        Inches(1), Inches(1), Inches(5), Inches(1)
    ).text_frame.text = "Electromagnetic induction generates EMF."

    path = tmp_path / "generator.pptx"
    presentation.save(path)

    result = read_pptx(path)

    assert "[SLIDE 1]" in result["text"]
    assert "DC Generator Construction" in result["text"]
    assert "[SLIDE 2]" in result["text"]
    assert "Working Principle" in result["text"]
    assert len(result["images"]) == 1
    assert result["images"][0]["slide_number"] == 1
    assert "armature and field poles" in result["images"][0]["context"]
    assert result["images"][0]["data"]


def test_read_pptx_rejects_wrong_extension(tmp_path):
    path = tmp_path / "not_a_presentation.txt"
    path.write_text("hello", encoding="utf-8")

    import pytest
    with pytest.raises(ValueError, match="\.pptx"):
        read_pptx(path)


def test_read_pptx_rejects_missing_file():
    import pytest
    with pytest.raises(FileNotFoundError):
        read_pptx("missing_presentation.pptx")
