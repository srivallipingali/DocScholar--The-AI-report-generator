from docx import Document
from docx.oxml.ns import qn

from report_generator.document_builder import build_structured_report
from report_generator.report_schema import Section, StructuredReport


def sample_config():
    return {
        "report_title": "",
        "font_name": "Times New Roman",
        "font_size": 12,
        "margins": {"top": 1, "bottom": 1, "left": 1, "right": 1},
    }


def sample_report():
    return StructuredReport(
        title="Solar Power",
        sections=[
            Section(heading="Introduction", content="Some text."),
            Section(heading="Conclusion", content="More text."),
        ],
    )


def instructions(document_part_element):
    """All field instruction strings (e.g. ' PAGE ') under an XML element."""
    return [
        node.text.strip()
        for node in document_part_element.iter(qn("w:instrText"))
    ]


def build(tmp_path, **options):
    path = tmp_path / "report.docx"
    build_structured_report(sample_report(), sample_config(), path, **options)
    return Document(path)


def test_toc_lists_every_heading_before_the_body(tmp_path):
    document = build(tmp_path, include_toc=True)
    texts = [p.text for p in document.paragraphs]

    assert "Table of Contents" in texts

    toc_start = texts.index("Table of Contents")
    # The entries appear right after the "Table of Contents" label ...
    assert texts[toc_start + 1 : toc_start + 3] == ["Introduction", "Conclusion"]
    # ... and the real headings come later, in the body.
    assert texts.count("Introduction") == 2


def test_toc_contains_a_toc_field(tmp_path):
    document = build(tmp_path, include_toc=True)

    assert any(
        text.startswith("TOC") for text in instructions(document.element.body)
    )


def test_toc_asks_word_to_update_fields_on_open(tmp_path):
    document = build(tmp_path, include_toc=True)

    update = document.settings.element.find(qn("w:updateFields"))

    assert update is not None
    assert update.get(qn("w:val")) == "true"


def test_no_toc_by_default(tmp_path):
    document = build(tmp_path)
    texts = [p.text for p in document.paragraphs]

    assert "Table of Contents" not in texts
    assert document.settings.element.find(qn("w:updateFields")) is None


def test_page_number_field_in_footer(tmp_path):
    document = build(tmp_path, page_numbers=True)
    footer = document.sections[0].footer._element

    assert "PAGE" in instructions(footer)


def test_no_page_numbers_by_default(tmp_path):
    document = build(tmp_path)
    footer = document.sections[0].footer._element

    assert instructions(footer) == []


def test_title_and_heading_use_template_font_not_theme_font(tmp_path):
    document = build(tmp_path)

    for style_name in ("Title", "Heading 1", "Normal"):
        fonts = document.styles[style_name].element.rPr.rFonts

        assert fonts.get(qn("w:ascii")) == "Times New Roman"
        assert fonts.get(qn("w:asciiTheme")) is None
