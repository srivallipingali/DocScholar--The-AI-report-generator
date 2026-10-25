
from docx import Document
import pytest

from report_generator.document_builder import build_report


def sample_config():
    return {
        "report_title": "Generated Report",
        "font_name": "Arial",
        "font_size": 12,
        "margins": {
            "top": 1,
            "bottom": 1,
            "left": 1,
            "right": 1
        }
    }


def test_build_report(tmp_path):
    output_path = tmp_path / "report.docx"

    text = (
        "Renewable energy comes from natural sources.\n\n"
        "Solar panels convert sunlight into electricity."
    )

    result = build_report(
        text,
        sample_config(),
        output_path
    )

    assert result == output_path
    assert output_path.is_file()

    document = Document(output_path)

    paragraphs = [
        paragraph.text for paragraph in document.paragraphs
    ]

    assert paragraphs == [
        "Generated Report",
        "Renewable energy comes from natural sources.",
        "Solar panels convert sunlight into electricity.",
    ]
    assert document.paragraphs[1].style.name != "Heading 1"


def test_build_report_requires_a_separate_title(tmp_path):
    config = sample_config()
    config["report_title"] = ""

    with pytest.raises(ValueError, match="report title is required"):
        build_report(
            "First line is report content.",
            config,
            tmp_path / "report.docx",
        )


def test_empty_text_is_rejected(tmp_path):
    with pytest.raises(ValueError, match="empty text"):
        build_report(
            "   ",
            sample_config(),
            tmp_path / "report.docx"
        )


# ---------- build_structured_report (the AI path) ----------

from io import BytesIO

from report_generator.document_builder import build_structured_report
from report_generator.report_schema import Section, StructuredReport


def sample_report():
    return StructuredReport(
        title="AI Chosen Title",
        sections=[
            Section(
                heading="Introduction",
                content="First paragraph.\n\nSecond paragraph.",
            ),
            Section(heading="Conclusion", content="The end."),
        ],
    )


def test_structured_report_contains_headings_and_paragraphs(tmp_path):
    output_path = tmp_path / "structured.docx"
    config = sample_config()
    config["report_title"] = ""  # no user title -> AI title is used

    build_structured_report(sample_report(), config, output_path)

    document = Document(output_path)
    paragraphs = [p.text for p in document.paragraphs]

    assert paragraphs == [
        "AI Chosen Title",
        "Introduction",
        "First paragraph.",
        "Second paragraph.",
        "Conclusion",
        "The end.",
    ]


def test_structured_report_user_title_overrides_ai_title(tmp_path):
    output_path = tmp_path / "structured.docx"

    build_structured_report(sample_report(), sample_config(), output_path)

    paragraphs = [p.text for p in Document(output_path).paragraphs]

    assert paragraphs[0] == "Generated Report"
    assert "AI Chosen Title" not in paragraphs


def test_structured_report_can_write_to_memory():
    buffer = BytesIO()

    build_structured_report(sample_report(), sample_config(), buffer)

    buffer.seek(0)
    assert len(Document(buffer).paragraphs) > 0


def test_structured_report_with_cover_page_and_index(tmp_path):
    output_path = tmp_path / "full_featured.docx"
    config = sample_config()
    config["report_title"] = ""
    config["include_cover_page"] = True
    config["include_index"] = True
    config["section_per_page"] = True
    config["author"] = "Research Team"
    config["template_name"] = "Academic Report"

    report = StructuredReport(
        title="Solar Power Future",
        subtitle="A Technical Evaluation",
        abstract="Executive overview of solar photovoltaics.",
        sections=[
            Section(heading="Introduction", content="Solar energy is growing."),
            Section(heading="Efficiency", content="Modern panels reach 24%."),
            Section(heading="Conclusion", content="Solar adoption is essential."),
        ],
    )

    build_structured_report(report, config, output_path)

    doc = Document(output_path)
    all_text = " ".join(p.text for p in doc.paragraphs)

    # Verify Cover Page items
    assert "ACADEMIC REPORT" in all_text
    assert "Solar Power Future" in all_text
    assert "A Technical Evaluation" in all_text
    assert "Executive overview of solar photovoltaics." in all_text
    assert "Prepared by: Research Team" in all_text

    # Verify Table of Contents
    assert "Table of Contents" in all_text
    assert "Introduction" in all_text
    assert "Efficiency" in all_text
    assert "Conclusion" in all_text

    # Verify page breaks (each topic on a new page)
    # In docx, page breaks are rendered as w:br w:type="page" or in runs
    page_break_count = sum(
        1 for p in doc.paragraphs
        for run in p.runs
        if "w:br" in run._r.xml and 'w:type="page"' in run._r.xml
    )
    # At least 4 page breaks: after cover, after index, after section 1, after section 2
    assert page_break_count >= 4


def test_structured_report_includes_uploaded_image_and_caption(tmp_path):
    from io import BytesIO
    from PIL import Image

    image_buffer = BytesIO()
    Image.new("RGB", (80, 50), color="red").save(image_buffer, format="PNG")

    output_path = tmp_path / "report_with_image.docx"
    build_structured_report(
        sample_report(),
        sample_config(),
        output_path,
        images=[
            {
                "name": "solar-panel.png",
                "data": image_buffer.getvalue(),
                "caption": "Solar panels in a field",
            }
        ],
    )

    document = Document(output_path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]

    assert "Uploaded Images" in paragraph_text
    assert "Figure 1. Solar panels in a field" in paragraph_text
    assert len(document.inline_shapes) == 1


def test_structured_report_uses_filename_when_image_caption_is_empty(tmp_path):
    from io import BytesIO
    from PIL import Image

    image_buffer = BytesIO()
    Image.new("RGB", (80, 50), color="blue").save(image_buffer, format="PNG")

    output_path = tmp_path / "report_with_default_caption.docx"
    build_structured_report(
        sample_report(),
        sample_config(),
        output_path,
        images=[
            {
                "name": "diagram.png",
                "data": image_buffer.getvalue(),
                "caption": "",
            }
        ],
    )

    document = Document(output_path)
    assert "Figure 1. diagram.png" in [
        paragraph.text for paragraph in document.paragraphs
    ]
    assert len(document.inline_shapes) == 1


def test_unselected_pptx_images_are_not_added_to_fallback_section(tmp_path):
    from io import BytesIO
    from PIL import Image

    image_buffer = BytesIO()
    Image.new("RGB", (80, 50), color="green").save(image_buffer, format="PNG")

    output_path = tmp_path / "report_selected_images_only.docx"
    build_structured_report(
        sample_report(),
        sample_config(),
        output_path,
        images=[{
            "id": "slide1:image_1",
            "source": "pptx",
            "name": "slide_1_image_1.png",
            "data": image_buffer.getvalue(),
            "caption": "A diagram",
        }],
    )

    document = Document(output_path)
    assert len(document.inline_shapes) == 0
    assert "Uploaded Images" not in [p.text for p in document.paragraphs]


def test_selected_image_is_placed_once_and_keeps_its_attribution(tmp_path):
    from io import BytesIO
    from PIL import Image

    image_buffer = BytesIO()
    Image.new("RGB", (80, 50), color="purple").save(image_buffer, format="PNG")
    image = {
        "id": "wikimedia:0",
        "source": "wikimedia",
        "name": "generator.png",
        "data": image_buffer.getvalue(),
        "caption": "A labeled diagram of a DC generator",
        "attribution": "File:Generator.png | Creator: Example | Source: Wikimedia Commons",
    }
    unplaced_image = {
        "id": "user-upload:1",
        "source": "upload",
        "name": "extra.png",
        "data": image_buffer.getvalue(),
        "caption": "An extra image",
    }
    report = sample_report()
    report.sections[0].image_ids = ["wikimedia:0"]
    report.sections[1].image_ids = ["wikimedia:0"]
    output_path = tmp_path / "report_with_attributed_image.docx"

    build_structured_report(
        report,
        sample_config(),
        output_path,
        images=[image, unplaced_image],
    )

    document = Document(output_path)
    paragraph_text = [paragraph.text for paragraph in document.paragraphs]
    assert len(document.inline_shapes) == 2
    assert "Figure 1. A labeled diagram of a DC generator" in paragraph_text
    assert "Figure 2. An extra image" in paragraph_text
    assert (
        "Source: File:Generator.png | Creator: Example | Source: Wikimedia Commons"
        in paragraph_text
    )
