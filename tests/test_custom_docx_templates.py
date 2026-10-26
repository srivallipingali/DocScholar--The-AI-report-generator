from io import BytesIO

from docx import Document

from report_generator.document_builder import build_report


def _template_bytes(paragraphs, header_text="TEMPLATE HEADER"):
    document = Document()
    document.sections[0].header.paragraphs[0].text = header_text
    for text in paragraphs:
        document.add_paragraph(text)
    stream = BytesIO()
    document.save(stream)
    return stream.getvalue()


def _paragraph_texts(path):
    return [paragraph.text for paragraph in Document(path).paragraphs]


def test_custom_docx_template_replaces_body_and_preserves_header(tmp_path):
    output = tmp_path / "report.docx"
    template = _template_bytes(["Old sample report", "Old content"])
    config = {
        "report_title": "New report",
        "include_cover_page": False,
        "font_name": "Arial",
        "font_size": 12,
        "margins": {"top": 1, "bottom": 1, "left": 1, "right": 1},
    }

    build_report(
        "Fresh generated content.",
        config,
        output,
        template_bytes=template,
        template_mode="replace_body",
    )

    result = Document(output)
    assert [p.text for p in result.paragraphs] == [
        "New report",
        "Fresh generated content.",
    ]
    assert result.sections[0].header.paragraphs[0].text == "TEMPLATE HEADER"


def test_custom_docx_template_fills_title_and_content_placeholders(tmp_path):
    output = tmp_path / "report.docx"
    template = _template_bytes([
        "[REPORT_TITLE]",
        "Introduction",
        "[REPORT_CONTENT]",
        "Prepared by the department",
    ])
    config = {
        "report_title": "Climate Report",
        "include_cover_page": False,
        "font_name": "Arial",
        "font_size": 12,
        "margins": {"top": 1, "bottom": 1, "left": 1, "right": 1},
    }

    build_report(
        "First paragraph.\n\nSecond paragraph.",
        config,
        output,
        template_bytes=template,
        template_mode="fill_placeholders",
    )

    assert _paragraph_texts(output) == [
        "Climate Report",
        "Introduction",
        "First paragraph.",
        "Second paragraph.",
        "Prepared by the department",
    ]


def test_placeholder_mode_reports_missing_content_placeholder(tmp_path):
    output = tmp_path / "report.docx"
    template = _template_bytes(["[REPORT_TITLE]", "No content marker here"])
    config = {
        "report_title": "Climate Report",
        "include_cover_page": False,
        "font_name": "Arial",
        "font_size": 12,
        "margins": {"top": 1, "bottom": 1, "left": 1, "right": 1},
    }

    import pytest

    with pytest.raises(ValueError, match="No content placeholder was found"):
        build_report(
            "Some content",
            config,
            output,
            template_bytes=template,
            template_mode="fill_placeholders",
        )
