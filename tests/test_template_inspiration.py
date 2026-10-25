import pymupdf
import pytest
from docx import Document
from docx.shared import Inches, Pt, RGBColor

from report_generator.template_inspiration import analyze_template


def test_analyze_docx_template_extracts_style_and_layout(tmp_path):
    file_path = tmp_path / "reference.docx"
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(30, 40, 50)
    title = document.styles["Title"]
    title.font.name = "Georgia"
    title.font.size = Pt(24)
    title.font.color.rgb = RGBColor(20, 60, 120)
    heading = document.styles["Heading 1"]
    heading.font.name = "Arial"
    heading.font.size = Pt(16)
    document.sections[0].left_margin = Inches(0.75)
    document.add_paragraph("A sample heading", style="Heading 1")
    document.add_paragraph("Template subject matter should not be included.")
    document.save(file_path)

    result = analyze_template(file_path)

    assert result["formatting"] == {
        "font_name": "Arial",
        "font_size": 10,
        "title_font_name": "Georgia",
        "title_font_size": 24,
        "heading_font_name": "Arial",
        "heading_font_size": 16,
        "accent_color": "#143C78",
        "text_color": "#1E2832",
        "margins": {"top": 1.0, "bottom": 1.0, "left": 0.75, "right": 1.25},
    }
    assert "Heading 1" in result["style_reference"]
    assert "Template subject matter" not in result["style_reference"]


def test_analyze_docx_template_without_title_or_heading_styles(tmp_path):
    file_path = tmp_path / "custom_styles.docx"
    document = Document()
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    for style_name in ("Title", "Heading 1"):
        style = document.styles[style_name]
        style._element.getparent().remove(style._element)
    document.add_paragraph("Plain template content")
    document.save(file_path)

    result = analyze_template(file_path)

    assert result["formatting"]["font_name"] == "Arial"
    assert result["formatting"]["font_size"] == 10
    assert "title_font_name" not in result["formatting"]
    assert "heading_font_name" not in result["formatting"]
    assert "No built-in section heading styles" in result["style_reference"]


def test_analyze_pdf_template_extracts_style_and_layout(tmp_path):
    file_path = tmp_path / "reference.pdf"
    document = pymupdf.open()
    page = document.new_page(width=612, height=792)
    page.insert_text((72, 72), "Sample report title", fontsize=24, color=(0, 0, 1))
    page.insert_text((72, 108), "Section heading", fontsize=16)
    page.insert_text((72, 140), "Sample body content", fontsize=11)
    document.save(file_path)
    document.close()

    result = analyze_template(file_path)

    assert result["formatting"]["title_font_size"] == 24
    assert result["formatting"]["heading_font_size"] == 16
    assert result["formatting"]["font_size"] == 11
    assert result["formatting"]["accent_color"] == "#0000FF"
    assert "body font" in result["style_reference"]
    assert "Sample body content" not in result["style_reference"]


def test_analyze_template_rejects_unsupported_files(tmp_path):
    file_path = tmp_path / "reference.txt"
    file_path.write_text("not a template", encoding="utf-8")

    with pytest.raises(ValueError, match="DOCX or PDF"):
        analyze_template(file_path)
