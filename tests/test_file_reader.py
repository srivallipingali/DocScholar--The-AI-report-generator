
import pytest
from docx import Document

from report_generator.file_reader import read_file


def test_read_txt_file(tmp_path):
    file_path = tmp_path / "sample.txt"
    file_path.write_text(
        "Renewable Energy\nClean energy matters.",
        encoding="utf-8"
    )

    result = read_file(file_path)

    assert result == (
        "Renewable Energy\nClean energy matters."
    )


def test_missing_file():
    with pytest.raises(FileNotFoundError):
        read_file("this_file_does_not_exist.txt")


def test_unsupported_file_type(tmp_path):
    file_path = tmp_path / "sample.csv"
    file_path.write_text("Test", encoding="utf-8")

    with pytest.raises(ValueError, match="Unsupported file type"):
        read_file(file_path)


def test_read_docx_file(tmp_path):
    file_path = tmp_path / "sample.docx"

    document = Document()
    document.add_paragraph("Renewable Energy")
    document.add_paragraph(
        "Solar panels convert sunlight into electricity."
    )

    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Energy Source"
    table.cell(0, 1).text = "Solar"

    document.save(file_path)

    result = read_file(file_path)

    assert "Renewable Energy" in result
    assert "Solar panels convert sunlight into electricity." in result
    assert "Energy Source" in result
    assert "Solar" in result


def test_empty_txt_file(tmp_path):
    file_path = tmp_path / "empty.txt"
    file_path.write_text("", encoding="utf-8")

    with pytest.raises(ValueError, match="no readable text"):
        read_file(file_path)


def test_empty_docx_file(tmp_path):
    file_path = tmp_path / "empty.docx"

    document = Document()
    document.save(file_path)

    with pytest.raises(ValueError, match="no readable text"):
        read_file(file_path)


