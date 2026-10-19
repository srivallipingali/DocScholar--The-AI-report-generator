
import json
from pathlib import Path
import pymupdf

from docx import Document
from docx.shared import Inches, Pt


def generate_report():
    # 1. Find the project directory
    project_directory = Path(__file__).resolve().parent

    # 2. Load the report formatting configuration
    config_path = (
        project_directory / "config" / "report_format.json"
    )

    with config_path.open("r", encoding="utf-8") as config_file:
        report_format = json.load(config_file)

    # 3. Ask the user for an input filename
    file_name = input(
        "Enter a TXT or PDF filename: "
    ).strip()

    # 4. Find and validate the input file
    source_path = project_directory / "samples" / file_name

    if not source_path.is_file():
        print("Error: File not found in the samples folder.")
        return

    file_extension = source_path.suffix.lower()

    # 5. Read the file according to its extension
    try:
        if file_extension == ".txt":
            source_content = source_path.read_text(
                encoding="utf-8"
            ).strip()

        elif file_extension == ".pdf":
            extracted_pages = []

            with pymupdf.open(source_path) as pdf:
                for page in pdf:
                    extracted_pages.append(page.get_text())

            source_content = "\n".join(extracted_pages).strip()

        else:
            print("Error: Only TXT and PDF files are supported.")
            return

    except (OSError, RuntimeError) as error:
        print(f"Error reading the input file: {error}")
        return

    # 6. Check whether readable text was extracted
    if not source_content:
        print(
            "Error: No readable text found in the file. "
            "The PDF may require OCR."
        )
        return

    # 7. Separate the title from the body
    lines = source_content.splitlines()

    source_title = lines[0].strip()
    body = "\n".join(lines[1:]).strip()

    # 8. Create a Word document
    document = Document()

    # 9. Apply the configured page margins
    page = document.sections[0]
    margins = report_format["margins"]

    page.top_margin = Inches(margins["top"])
    page.bottom_margin = Inches(margins["bottom"])
    page.left_margin = Inches(margins["left"])
    page.right_margin = Inches(margins["right"])

    # 10. Apply the configured default font
    normal_style = document.styles["Normal"]

    normal_style.font.name = report_format["font_name"]
    normal_style.font.size = Pt(report_format["font_size"])

    # 11. Add the report title and source heading
    document.add_heading(
        report_format["report_title"],
        level=0
    )

    document.add_heading(source_title, level=1)

    # 12. Add the source text as paragraphs
    for paragraph in body.split("\n\n"):
        paragraph = " ".join(paragraph.split())

        if paragraph:
            document.add_paragraph(paragraph)

    # 13. Create the output folder if necessary
    output_directory = project_directory / "output"
    output_directory.mkdir(parents=True, exist_ok=True)

    # 14. Generate an output filename from the input filename
    output_filename = f"{source_path.stem}_report.docx"
    output_path = output_directory / output_filename

    # 15. Save the Word report
    try:
        document.save(output_path)
    except OSError as error:
        print(f"Error saving the report: {error}")
        return

    print(f"\nReport generated successfully!")
    print(f"Output file: {output_path}")


if __name__ == "__main__":
    generate_report()

